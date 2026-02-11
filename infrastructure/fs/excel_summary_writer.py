# infrastructure/fs/excel_summary_writer.py
from __future__ import annotations

from pathlib import Path
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


class ExcelSummaryWriter:
    def write_residuos(
        self,
        xlsx_path: Path,
        rows: List[Dict[str, Any]],
        errors: List[Dict[str, Any]],
    ) -> None:
        xlsx_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "residuos"

        headers = [
            "archivo",
            "proveedor",
            "obra",
            "planta_tipo",
            "ler",
            "tipo_residuo",
            "cantidad",
            "unidad",
            "confianza_pct",
            "observaciones",
        ]
        ws.append(headers)

        for r in rows:
            ws.append([r.get(h) for h in headers])

        if len(rows) > 0:
            last_row = len(rows) + 1
            last_col = len(headers)
            table = Table(displayName="TablaResiduos", ref=f"A1:{self._col(last_col)}{last_row}")
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            ws.add_table(table)
            ws.auto_filter.ref = table.ref

        # Sheet errores
        ws_err = wb.create_sheet("errores")
        ws_err.append(["archivo", "error"])
        for e in errors:
            ws_err.append([e.get("archivo"), e.get("error")])

        self._autosize(ws)
        self._autosize(ws_err)

        wb.save(xlsx_path)

    @staticmethod
    def _col(n: int) -> str:
        # 1->A, 26->Z, 27->AA
        s = ""
        while n:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s

    @staticmethod
    def _autosize(ws) -> None:
        for col_cells in ws.columns:
            max_len = 0
            col = col_cells[0].column_letter
            for c in col_cells:
                if c.value is None:
                    continue
                max_len = max(max_len, len(str(c.value)))
            ws.column_dimensions[col].width = min(max(10, max_len + 2), 60)
