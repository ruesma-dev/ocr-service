# simulation/service3_residuos_excel_exporter.py
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


class Service3ResiduosExcelExporter:
    """
    SIMULACIÓN DEL SERVICIO 3 (persistencia/reporting).
    - Se alimenta de los envelopes JSON producidos por el extractor.
    - Genera un Excel resumen.
    - Este módulo es intencionalmente aislado para poder borrarlo fácilmente.
    """

    def export(
        self,
        *,
        xlsx_path: Path,
        envelopes: List[Dict[str, Any]],
        errors: List[Dict[str, Any]],
    ) -> None:
        xlsx_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "residuos"

        headers = [
            "archivo",
            "proveedor",
            "fecha_documento",
            "periodo_inicio",
            "periodo_fin",
            "obra",
            "planta_tipo",
            "ler",
            "tipo_residuo",
            "cantidad",
            "unidad",
            "confianza_pct",
            "observaciones",
            "prompt_key",
            "schema",
            "model",
            "sha256",
        ]
        ws.append(headers)

        rows_count = 0

        for env in envelopes:
            meta = env.get("meta") or {}
            data = env.get("data") or {}

            residuos = data.get("residuos") or []
            if not isinstance(residuos, list):
                residuos = []

            base = {
                "archivo": meta.get("source_filename"),
                "proveedor": data.get("proveedor"),
                "fecha_documento": data.get("fecha_documento"),
                "periodo_inicio": data.get("periodo_inicio"),
                "periodo_fin": data.get("periodo_fin"),
                "obra": data.get("obra"),
                "planta_tipo": data.get("planta_tipo"),
                "confianza_pct": data.get("confianza_pct"),
                "observaciones": data.get("observaciones"),
                "prompt_key": meta.get("prompt_key"),
                "schema": meta.get("schema"),
                "model": meta.get("model"),
                "sha256": meta.get("source_sha256"),
            }

            # 1 fila por residuo
            if residuos:
                for item in residuos:
                    ws.append(
                        [
                            base["archivo"],
                            base["proveedor"],
                            base["fecha_documento"],
                            base["periodo_inicio"],
                            base["periodo_fin"],
                            base["obra"],
                            base["planta_tipo"],
                            (item or {}).get("ler"),
                            (item or {}).get("tipo_residuo"),
                            (item or {}).get("cantidad"),
                            (item or {}).get("unidad"),
                            base["confianza_pct"],
                            base["observaciones"],
                            base["prompt_key"],
                            base["schema"],
                            base["model"],
                            base["sha256"],
                        ]
                    )
                    rows_count += 1
            else:
                # fila vacía para trazabilidad del documento
                ws.append(
                    [
                        base["archivo"],
                        base["proveedor"],
                        base["fecha_documento"],
                        base["periodo_inicio"],
                        base["periodo_fin"],
                        base["obra"],
                        base["planta_tipo"],
                        None,
                        None,
                        None,
                        None,
                        base["confianza_pct"],
                        base["observaciones"],
                        base["prompt_key"],
                        base["schema"],
                        base["model"],
                        base["sha256"],
                    ]
                )
                rows_count += 1

        if rows_count > 0:
            last_row = rows_count + 1
            last_col = len(headers)
            table = Table(displayName="TablaResiduos", ref=f"A1:{self._col(last_col)}{last_row}")
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            ws.add_table(table)
            ws.auto_filter.ref = table.ref

        # Hoja errores
        ws_err = wb.create_sheet("errores")
        ws_err.append(["archivo", "error"])
        for e in errors:
            ws_err.append([e.get("archivo"), e.get("error")])

        self._autosize(ws)
        self._autosize(ws_err)

        wb.save(xlsx_path)

    @staticmethod
    def _col(n: int) -> str:
        s = ""
        while n:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s

    @staticmethod
    def _autosize(ws) -> None:
        for col_cells in ws.columns:
            max_len = 0
            col = col_cells[0].column_letter
            for c in col_cells:
                if c.value is None:
                    continue
                max_len = max(max_len, len(str(c.value)))
            ws.column_dimensions[col].width = min(max(10, max_len + 2), 70)
